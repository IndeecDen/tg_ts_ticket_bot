import asyncio
import ast
import html
import shutil
import sqlite3
import tempfile
from contextlib import closing
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

from aiogram.exceptions import TelegramBadRequest, TelegramRetryAfter
from aiogram.methods import SendMessage
from db.database import Database
from db.async_database import AsyncDatabase
from services.requests import RequestService, request_card
from utils.text import plain_text, split_text, utf16_size
from handlers import commands, callbacks_stats, callbacks_requests
from config import Config
from handlers import common

ROOT = Path(__file__).resolve().parents[1]


class DatabaseFixture(unittest.TestCase):
    def setUp(self):
        base = ROOT / '.test-data'
        base.mkdir(exist_ok=True)
        self.directory = Path(tempfile.mkdtemp(dir=base))
        self.path = self.directory / 'test.db'
        self.db = Database(str(self.path))

    def tearDown(self):
        # Only remove the isolated directory created by this test.
        self.assertTrue(self.directory.resolve().is_relative_to((ROOT / '.test-data').resolve()))
        shutil.rmtree(self.directory)

    def collect(self, user=1, chat=10, text='problem', **extra):
        self.db.collect_message(chat, user, 'client', {'message_id': 42, 'text': text, **extra}, 0)


class DatabaseTests(DatabaseFixture):
    def test_pending_survives_restart_and_preserves_full_payload(self):
        self.collect(text='long text' * 1000, photo='file-id')
        restarted = Database(str(self.path))
        request_id, = restarted.promote_due_requests()
        self.assertEqual(restarted.get_request_by_id(request_id)[3], 'long text' * 1000)
        self.assertEqual(restarted.delivery_payload(request_id)[0]['photo'], 'file-id')
        self.assertEqual(restarted.promote_due_requests(), [])
        self.assertEqual(len(restarted.due_deliveries()), 2)

    def test_duplicate_messages_and_concurrent_promotion(self):
        self.collect()
        self.collect()
        from concurrent.futures import ThreadPoolExecutor
        with ThreadPoolExecutor(2) as pool:
            results = list(pool.map(lambda _: self.db.promote_due_requests(), range(2)))
        self.assertEqual(sum(map(len, results)), 1)
        with self.db.connect() as conn:
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM requests').fetchone()[0], 1)
        self.assertEqual(len(self.db.delivery_payload(1)), 1)

    def test_two_clients_in_same_chat_are_collected(self):
        self.collect(user=1)
        self.collect(user=2)
        self.assertEqual(len(self.db.promote_due_requests()), 2)

    def test_staff_response_cancels_pending_only_in_its_chat(self):
        self.collect(chat=10)
        self.collect(chat=20)
        self.assertEqual(self.db.cancel_pending(10), 1)
        request_id, = self.db.promote_due_requests()
        self.assertEqual(self.db.get_request_by_id(request_id)[8], 20)

    def test_cancel_before_timeout(self):
        self.collect()
        self.assertTrue(self.db.cancel_request(1, 10))
        self.assertEqual(self.db.promote_due_requests(), [])

    def test_identity_is_id_not_name(self):
        rid = self.db.create_request(1, 'client', 'text', 10, 42)
        self.assertTrue(self.db.assign_specialist_if_new(rid, 'Same Name', 101))
        self.assertFalse(self.db.close_request(rid, 202))
        self.assertEqual(len(self.db.get_active_requests_by_specialist(101)), 1)
        self.assertEqual(self.db.get_active_requests_by_specialist(202), [])
        self.assertTrue(self.db.close_request(rid, 101))
        self.assertEqual(len(self.db.get_specialist_statistics('day', specialist_id=101)), 1)
        self.assertEqual(self.db.get_specialist_statistics('day', specialist_id=202), [])

    def test_duplicate_names_have_separate_statistics(self):
        for sid in (101, 202):
            rid = self.db.create_request(sid, 'client', 'text', 10, sid)
            self.db.assign_specialist_if_new(rid, 'Same Name', sid)
            self.db.close_request(rid, sid)
        stats = self.db.get_specialist_statistics('day')
        self.assertEqual(len(stats), 2)
        self.assertNotEqual(stats[0][0], stats[1][0])

    def test_legacy_migration_preserves_settings_and_backs_up(self):
        old_path = self.directory / 'old.db'
        with closing(sqlite3.connect(old_path)) as conn:
            conn.execute('CREATE TABLE unassigned_reminder_settings (id INTEGER PRIMARY KEY, enabled INTEGER, weekdays TEXT, threshold_minutes INTEGER)')
            conn.execute("INSERT INTO unassigned_reminder_settings VALUES (1, 1, '2,4', 45)")
            conn.commit()
        upgraded = Database(str(old_path))
        settings = upgraded.get_unassigned_reminder_settings()
        self.assertTrue(settings['enabled'])
        self.assertEqual(settings['weekdays'], [2, 4])
        self.assertEqual(settings['threshold_minutes'], 45)
        self.assertEqual(settings['interval_minutes'], 60)
        self.assertEqual(len(list(self.directory.glob('old.db.backup-*'))), 1)
        Database(str(old_path))
        self.assertEqual(len(list(self.directory.glob('old.db.backup-*'))), 1)

    def test_old_undelivered_request_is_requeued(self):
        self.db.create_request(1, 'client', 'text', 10, 42)
        restarted = Database(str(self.path))
        self.assertEqual(len(restarted.due_deliveries()), 2)

    def test_migration_on_readonly_copy_of_existing_database(self):
        original = ROOT / 'bot.db'
        if not original.exists():
            self.skipTest('No existing local database')
        clone = self.directory / 'existing-copy.db'
        with closing(sqlite3.connect(original.as_uri() + '?mode=ro', uri=True)) as source:
            before = source.execute('SELECT * FROM requests ORDER BY id').fetchall()
            with closing(sqlite3.connect(clone)) as target:
                source.backup(target)
        migrated = Database(str(clone))
        with migrated.connect() as conn:
            self.assertEqual(conn.execute('SELECT * FROM requests ORDER BY id').fetchall(), before)
            self.assertEqual(conn.execute('PRAGMA quick_check').fetchone()[0], 'ok')
            self.assertGreaterEqual(len(conn.execute('PRAGMA index_list(requests)').fetchall()), 4)

    def test_stale_delivery_cannot_remove_new_status_notification(self):
        self.collect()
        rid, = self.db.promote_due_requests()
        revision = self.db.due_deliveries()[0][2]
        self.db.assign_specialist_if_new(rid, 'Specialist', 101)
        self.db.delivery_succeeded(rid, 'work', revision, 99)
        self.assertEqual(self.db.get_request_by_id(rid)[11], 99)
        self.assertEqual(len(self.db.due_deliveries()), 2)

    def test_cleanup_retains_failed_messages_and_skips_delivery_queue(self):
        rid = self.db.create_request(1, 'client', 'text', 10, 42)
        self.db.update_work_message_id(rid, 100)
        self.db.update_bot_message_id(rid, 200)
        self.db.close_all_requests()
        self.assertEqual(self.db.get_all_bot_messages(), [])
        for req, dest, rev, _ in self.db.due_deliveries():
            self.db.delivery_succeeded(req, dest, rev, 100 if dest == 'work' else 200)
        self.db.mark_message_deleted(10, 200, -100999)
        self.assertEqual(self.db.get_all_bot_messages(), [(10, None, 100)])


class AsyncTests(DatabaseFixture, unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.adb = AsyncDatabase(self.db)
        self.bot = SimpleNamespace(
            send_message=AsyncMock(return_value=SimpleNamespace(message_id=100)),
            send_photo=AsyncMock(return_value=SimpleNamespace(message_id=101)),
            edit_message_text=AsyncMock(), edit_message_caption=AsyncMock())
        self.config = SimpleNamespace(work_chat_id=-100999, response_timeout=0)
        self.service = RequestService(self.bot, self.adb, self.config)

    async def test_failed_send_is_retried_after_restart_without_new_request(self):
        self.collect()
        async def fail_work(**kwargs):
            if kwargs['chat_id'] == self.config.work_chat_id:
                raise OSError('offline')
            return SimpleNamespace(message_id=200)
        self.bot.send_message.side_effect = fail_work
        await self.service.tick()
        self.assertEqual(self.db.get_request_by_id(1)[11], None)
        self.assertEqual(self.db.get_request_by_id(1)[10], 200)
        with self.db.connect() as conn:
            conn.execute("UPDATE delivery_queue SET next_attempt=''")
        self.bot.send_message.side_effect = None
        restarted = RequestService(self.bot, AsyncDatabase(Database(str(self.path))), self.config)
        await restarted.tick()
        self.assertEqual(self.db.due_deliveries(), [])
        self.assertEqual(len(self.db.get_open_requests()), 1)
        self.assertEqual(self.bot.send_message.await_count, 3)

    async def test_photo_caption_updates_fit_and_escape(self):
        self.collect(text='<&😀>' * 1500, photo='file-id')
        await self.service.tick()
        self.db.assign_specialist_if_new(1, 'Name <&>', 101)
        await self.service.tick()
        self.db.close_request(1, 101)
        await self.service.tick()
        for call in self.bot.edit_message_caption.await_args_list:
            caption = call.kwargs['caption']
            self.assertLessEqual(utf16_size(caption), 1024)
            self.assertIn('Name &lt;&amp;&gt;', caption)
        self.assertEqual(self.bot.edit_message_caption.await_count, 2)

    async def test_slow_chat_does_not_hold_other_chats(self):
        self.collect(user=1, chat=10)
        self.collect(user=2, chat=20)
        other_sent = asyncio.Event()
        async def send(**kwargs):
            if '#1' in kwargs['text']:
                await asyncio.wait_for(other_sent.wait(), 2)
            else:
                other_sent.set()
            return SimpleNamespace(message_id=100)
        self.bot.send_message.side_effect = send
        await asyncio.wait_for(self.service.tick(), 3)
        self.assertTrue(other_sent.is_set())
        self.assertEqual(self.db.due_deliveries(), [])

    async def test_retry_after_is_persisted(self):
        self.collect()
        self.bot.send_message.side_effect = TelegramRetryAfter(method=SendMessage(chat_id=1, text='x'), message='flood', retry_after=60)
        await self.service.tick()
        self.assertEqual(self.db.due_deliveries(), [])
        with self.db.connect() as conn:
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM delivery_queue').fetchone()[0], 2)

    async def test_unauthorized_stats_callback_is_rejected(self):
        callback = SimpleNamespace(message=SimpleNamespace(chat=SimpleNamespace(id=10, type='group')),
            from_user=SimpleNamespace(id=202), bot=self.bot, data='stats:excel:day:specialists', answer=AsyncMock())
        state = AsyncMock()
        with patch.object(callbacks_stats, 'check_user_role', AsyncMock(return_value={'is_admin_or_botadmin': False, 'is_specialist': False})), \
             patch.object(callbacks_stats, '_stats_output', AsyncMock()) as output:
            await callbacks_stats.cb_stats(callback, state)
            output.assert_not_awaited()
            self.assertTrue(callback.answer.await_args.kwargs['show_alert'])

    async def test_done_rejects_other_specialist_with_same_name(self):
        self.collect()
        self.db.promote_due_requests()
        self.db.assign_specialist_if_new(1, 'Same Name', 101)
        callback = SimpleNamespace(message=SimpleNamespace(chat=SimpleNamespace(id=-100999, type='supergroup')),
            from_user=SimpleNamespace(id=202, full_name='Same Name'), bot=self.bot, data='done:1', answer=AsyncMock())
        with patch.object(callbacks_requests, 'database', self.adb), \
             patch.object(callbacks_requests, 'check_user_role', AsyncMock(return_value={'is_admin_or_botadmin': False, 'is_specialist': True})):
            await callbacks_requests.cb_done(callback, -100999, self.service)
        self.assertEqual(self.db.get_request_by_id(1)[4], 'in_progress')

    async def test_sqlite_runs_outside_event_loop(self):
        import threading
        db = AsyncDatabase(SimpleNamespace(thread_id=threading.get_ident))
        self.assertNotEqual(await db.thread_id(), threading.get_ident())

    async def test_real_dispatcher_routes_admin_client_and_fsm(self):
        from aiogram import Bot
        from aiogram.client.session.base import BaseSession
        from aiogram.types import Update, Message, Chat, User, ChatMemberLeft
        from datetime import datetime
        from main import build_dispatcher

        class OfflineSession(BaseSession):
            def __init__(self):
                super().__init__()
                self.calls = []

            async def make_request(self, bot, method, timeout=None):
                self.calls.append(method)
                if method.__api_method__ == 'getChatMember':
                    return ChatMemberLeft(user=User(id=method.user_id, is_bot=False, first_name='User'))
                if method.__api_method__ in ('sendMessage', 'editMessageText'):
                    return Message(message_id=len(self.calls), date=datetime.now(),
                        chat=Chat(id=method.chat_id, type='private'), text=method.text)
                return True

            async def close(self):
                pass

            async def stream_content(self, *args, **kwargs):
                yield b''

        session = OfflineSession()
        bot = Bot('123456:offline-test-token', session=session)
        service = RequestService(bot, self.adb, self.config)
        self.db.add_botadmin(42, 'admin')
        dp = build_dispatcher(self.adb, self.config, service)

        async def send(user_id, text):
            await dp.feed_update(bot, Update(update_id=len(session.calls)+1,
                message=Message(message_id=len(session.calls)+1, date=datetime.now(),
                    chat=Chat(id=user_id, type='private'),
                    from_user=User(id=user_id, is_bot=False, first_name='User'), text=text)))

        await send(42, '/get_daily_reminder')
        await send(42, '/set_notice')
        await send(42, 'Maintenance tomorrow')
        await send(42, '-')
        await send(42, '-')
        self.assertEqual(len(self.db.get_all_announcements()), 1)
        await send(1, 'My problem')
        await service.tick()
        self.assertEqual(len(self.db.get_open_requests()), 1)
        await send(1, '/cancel')
        self.assertEqual(self.db.get_request_by_id(1)[4], 'canceled')
        await send(42, '/set_notice')
        await send(42, '/cancel')
        self.assertEqual(len(self.db.get_all_announcements()), 1)
        await dp.storage.close()
        await bot.session.close()

    async def test_owner_can_close_after_renaming(self):
        self.collect()
        self.db.promote_due_requests()
        self.db.assign_specialist_if_new(1, 'Old Name', 101)
        callback = SimpleNamespace(message=SimpleNamespace(chat=SimpleNamespace(id=-100999, type='supergroup')),
            from_user=SimpleNamespace(id=101, full_name='New Name'), bot=self.bot, data='done:1', answer=AsyncMock())
        with patch.object(callbacks_requests, 'database', self.adb), \
             patch.object(callbacks_requests, 'check_user_role', AsyncMock(return_value={'is_admin_or_botadmin': False, 'is_specialist': True})):
            await callbacks_requests.cb_done(callback, -100999, self.service)
        self.assertEqual(self.db.get_request_by_id(1)[4], 'closed')

    async def test_long_reports_are_sent_in_multiple_messages(self):
        from utils.telegram_utils import safe_send_message, safe_edit_message_text
        content = '<b>' + html.escape('<&😀>' * 2000) + '</b>'
        await safe_send_message(self.bot, 1, content)
        self.assertGreater(self.bot.send_message.await_count, 1)
        for call in self.bot.send_message.await_args_list:
            self.assertLessEqual(utf16_size(call.kwargs['text']), 4000)
        await safe_edit_message_text(self.bot, 1, 100, content)
        self.bot.edit_message_text.assert_awaited_once()

    async def test_cancellation_keeps_delivery_for_restart(self):
        self.collect()
        self.bot.send_message.side_effect = asyncio.CancelledError()
        with self.assertRaises(asyncio.CancelledError):
            await self.service.tick()
        self.assertEqual(len(self.db.due_deliveries()), 2)


class TextAndConfigTests(unittest.TestCase):
    def test_command_placeholders_are_not_interpreted_as_html(self):
        text = 'Usage: /assign_request <request_id> <user_id>'
        self.assertEqual(plain_text(''.join(split_text(text))), text)

    def test_excel_treats_user_formulas_as_text(self):
        from io import BytesIO
        import openpyxl
        from utils.excel import save_workbook
        workbook = openpyxl.Workbook()
        workbook.active.append(['=1+1', 'ordinary text'])
        output = BytesIO()
        save_workbook(workbook, output)
        output.seek(0)
        loaded = openpyxl.load_workbook(output)
        self.assertEqual(loaded.active['A1'].data_type, 's')
        self.assertEqual(loaded.active['A1'].value, '=1+1')

    def test_long_escaped_names_fit_caption(self):
        request = (1, 2, '&' * 100, '<>' * 1000, 'in_progress', '&' * 100,
                   '2026-01-01T10:00:00', None, -100999, 42, None, None,
                   '2026-01-01T09:00:00', 101)
        text, _ = request_card(request, 'work')
        self.assertLessEqual(utf16_size(text), 1024)

    def test_html_split_preserves_content_and_balances_tags(self):
        import xml.etree.ElementTree as ET
        original = '<&😀>' * 2500
        text = '<b>Header</b>\n<a href="tg://user?id=123">' + html.escape(original) + '</a>'
        chunks = split_text(text)
        self.assertGreater(len(chunks), 1)
        for chunk in chunks:
            self.assertLessEqual(utf16_size(chunk), 4000)
            ET.fromstring('<root>' + chunk + '</root>')
        self.assertEqual(''.join(plain_text(c) for c in chunks), 'Header\n' + original)

    def test_plain_split_preserves_emoji(self):
        text = '😀' * 5000
        chunks = split_text(text, None)
        self.assertEqual(''.join(chunks), text)
        self.assertTrue(all(utf16_size(c) <= 4000 for c in chunks))

    def test_config_does_not_change_memory_when_save_fails(self):
        config = Config(_env_file=None, bot_token='test', work_chat_id=-100999)
        with patch('config.set_key', side_effect=OSError('read only')):
            self.assertFalse(config.update_response_timeout(12))
        self.assertEqual(config.response_timeout, 300)

    def test_sources_compile(self):
        files = [ROOT/'main.py', ROOT/'config.py']
        for folder in ('db', 'handlers', 'services', 'utils', 'tests'):
            files.extend((ROOT/folder).glob('*.py'))
        for path in files:
            ast.parse(path.read_text(encoding='utf-8'), filename=str(path))


if __name__ == '__main__':
    unittest.main()
