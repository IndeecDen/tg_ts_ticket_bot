"""
Обработчики callback-кнопок статистики: stats:, my_stats:.
"""
from aiogram import Router, F, types
from aiogram.types import CallbackQuery, InlineKeyboardMarkup, InlineKeyboardButton
from aiogram.fsm.context import FSMContext
from db import database
from utils.telegram_utils import safe_edit_message_text, safe_answer_callback
from utils.logging_config import logger
from .commands import StatsStates, check_user_role
import html
import asyncio
import openpyxl
from utils.excel import save_workbook
from openpyxl.utils import get_column_letter
from tempfile import NamedTemporaryFile
import os
from datetime import datetime, timedelta
from typing import Optional

router = Router()

def _valid_stats_data(data):
    parts = data.split(":")
    personal = parts[0] == "my_stats"
    if len(parts) == 2:
        return not personal and parts[1] in ("specialists", "clients")
    if len(parts) < 3 or parts[1] not in ("period", "telegram", "excel"):
        return False
    if parts[2] not in ("day", "week", "month", "custom"):
        return False
    custom_output = parts[2] == "custom" and parts[1] != "period"
    expected = (5 if personal else 6) if custom_output else (3 if personal else 4)
    if len(parts) != expected:
        return False
    if not personal and parts[-1] not in ("specialists", "clients"):
        return False
    if custom_output:
        try:
            start, end = (datetime.strptime(value, "%Y-%m-%d") for value in parts[3:5])
            return start <= end
        except ValueError:
            return False
    return True




# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _fmt_td(seconds: float) -> str:
    h = int(seconds // 3600)
    m = int((seconds % 3600) // 60)
    s = int(seconds % 60)
    return f"{h:02d}:{m:02d}:{s:02d}"


def format_period(period: str, start_date: Optional[str] = None, end_date: Optional[str] = None) -> str:
    today = datetime.now().date()
    if period == "custom" and start_date and end_date:
        return f"Период: {start_date} - {end_date}"
    elif period == "day":
        s = today.strftime('%Y-%m-%d')
        return f"Период: {s} - {s}"
    elif period == "week":
        s = (today - timedelta(days=7)).strftime('%Y-%m-%d')
        return f"Период: {s} - {today.strftime('%Y-%m-%d')}"
    elif period == "month":
        s = (today - timedelta(days=30)).strftime('%Y-%m-%d')
        return f"Период: {s} - {today.strftime('%Y-%m-%d')}"
    return "Период: Не указан"


# ---------------------------------------------------------------------------
# stats: — общая статистика
# ---------------------------------------------------------------------------

@router.callback_query(F.data.startswith("stats:"))
async def cb_stats(callback: CallbackQuery, state: FSMContext) -> None:
    if not callback.message:
        await safe_answer_callback(callback, "Сообщение недоступно", show_alert=True)
        return
    roles = await check_user_role(callback.bot, callback.from_user.id, callback.message.chat.id, callback.message.chat.type)
    allowed = roles["is_admin_or_botadmin"]
    if not allowed:
        await state.clear()
        await safe_answer_callback(callback, "⛔ Доступ запрещён", show_alert=True)
        return
    if not _valid_stats_data(callback.data):
        await safe_answer_callback(callback, "Некорректная кнопка статистики", show_alert=True)
        return
    parts = callback.data.split(":")
    action = parts[1]

    if action in ("specialists", "clients"):
        await state.update_data(stats_type=action)
        keyboard = InlineKeyboardMarkup(inline_keyboard=[
            [
                InlineKeyboardButton(text="День", callback_data=f"stats:period:day:{action}"),
                InlineKeyboardButton(text="Неделя", callback_data=f"stats:period:week:{action}")
            ],
            [
                InlineKeyboardButton(text="Месяц", callback_data=f"stats:period:month:{action}"),
                InlineKeyboardButton(text="Произвольный период", callback_data=f"stats:period:custom:{action}")
            ]
        ])
        await safe_edit_message_text(
            callback.bot,
            chat_id=callback.message.chat.id,
            message_id=callback.message.message_id,
            text="📊 Выберите временной интервал для статистики:",
            reply_markup=keyboard
        )
        await safe_answer_callback(callback, "")

    elif action == "period":
        period = parts[2]
        stats_type = parts[3]
        if period == "custom":
            await state.update_data(stats_type=parts[3])
            await safe_edit_message_text(
                callback.bot,
                chat_id=callback.message.chat.id,
                message_id=callback.message.message_id,
                text="📅 Введите начальную дату в формате ГГГГ-ММ-ДД (например, 2025-05-01):",
                reply_markup=None
            )
            await state.set_state(StatsStates.waiting_for_custom_start_date)
        else:
            keyboard = InlineKeyboardMarkup(inline_keyboard=[[
                InlineKeyboardButton(text="В Telegram", callback_data=f"stats:telegram:{period}:{stats_type}"),
                InlineKeyboardButton(text="Скачать Excel", callback_data=f"stats:excel:{period}:{stats_type}")
            ]])
            await safe_edit_message_text(
                callback.bot,
                chat_id=callback.message.chat.id,
                message_id=callback.message.message_id,
                text="📊 Выберите формат вывода статистики:",
                reply_markup=keyboard
            )
        await safe_answer_callback(callback, "")

    elif action in ("telegram", "excel"):
        await _stats_output(callback, parts)


async def _stats_output(callback: CallbackQuery, parts: list) -> None:
    output_format = parts[1]
    period = parts[2]
    stats_type = parts[5] if len(parts) > 5 else parts[3] if len(parts) > 3 else None
    start_date = parts[3] if len(parts) > 5 else None
    end_date = parts[4] if len(parts) > 5 else None

    if stats_type not in ("specialists", "clients"):
        logger.error(f"Некорректный stats_type: {stats_type}, parts: {parts}")
        await safe_answer_callback(callback, "Ошибка при обработке статистики. Попробуйте снова.", show_alert=True)
        return

    if stats_type == "specialists":
        stats = (await database.get_specialist_statistics(period, start_date, end_date))
        title = "📊 Статистика по специалистам"
        entity_label = "👨‍💻"
        excel_filename = f"specialist_stats_{period}"
        excel_sheet_title = "Specialist Statistics"
        headers = ["Специалист", "Клиенты", "Количество заявок", "Время исполнения"]
    else:
        stats = (await database.get_client_chat_statistics(period, start_date, end_date))
        title = "📊 Статистика по клиентским чатам"
        entity_label = "💬"
        excel_filename = f"client_chat_stats_{period}"
        excel_sheet_title = "Client Chat Statistics"
        headers = ["Название чата", "Исполнители", "Количество заявок", "Общее время исполнения (HH:MM:SS)", "Общее время ожидания (HH:MM:SS)"]

    if not stats:
        period_text = format_period(period, start_date, end_date)
        await safe_edit_message_text(
            callback.bot,
            chat_id=callback.message.chat.id,
            message_id=callback.message.message_id,
            text=f"{title} ({period_text})\n\nНет данных для выбранного периода.",
            reply_markup=None
        )
        await safe_answer_callback(callback, "Нет данных для отображения")
        return

    if stats_type == "clients":
        chat_stats: dict = {}
        for chat_id, _, count, exec_time, wait_time in stats:
            if chat_id not in chat_stats:
                chat_stats[chat_id] = {"total_requests": 0, "total_execution_time": 0.0, "total_waiting_time": 0.0}
            chat_stats[chat_id]["total_requests"] += int(count)
            chat_stats[chat_id]["total_execution_time"] += exec_time or 0
            chat_stats[chat_id]["total_waiting_time"] += wait_time or 0
        total_requests = sum(int(c) for _, _, c, _, _ in stats)
        total_exec = sum(t for _, _, _, t, _ in stats if t)
        total_wait = sum(w for _, _, _, _, w in stats if w)
    else:
        spec_stats: dict = {}
        for specialist, _, count, exec_time in stats:
            if specialist not in spec_stats:
                spec_stats[specialist] = {"total_requests": 0, "total_execution_time": 0.0}
            spec_stats[specialist]["total_requests"] += int(count)
            spec_stats[specialist]["total_execution_time"] += exec_time or 0
        total_requests = sum(int(c) for _, _, c, _ in stats)
        total_exec = sum(t for _, _, _, t in stats if t)
        total_wait = None

    total_exec_str = _fmt_td(total_exec) if total_exec else "н/д"
    total_wait_str = _fmt_td(total_wait) if total_wait and stats_type == "clients" else None

    if output_format == "telegram":
        period_text = format_period(period, start_date, end_date)
        response = f"{title} ({period_text})\n\n"
        if stats_type == "clients":
            for chat_id, data in chat_stats.items():
                try:
                    chat = await callback.bot.get_chat(int(chat_id))
                    name = html.escape(chat.title or f"Чат {chat_id}")
                except Exception:
                    name = f"Чат {chat_id}"
                e_str = _fmt_td(data["total_execution_time"]) if data["total_execution_time"] else "н/д"
                w_str = _fmt_td(data["total_waiting_time"]) if data["total_waiting_time"] else "н/д"
                response += f"{entity_label} {name}: {data['total_requests']} заявок, исполнение {e_str}, ожидание {w_str}\n"
            response += f"\nВсего: {total_requests} заявок, исполнение {total_exec_str}, ожидание {total_wait_str}"
        else:
            for specialist, data in spec_stats.items():
                e_str = _fmt_td(data["total_execution_time"]) if data["total_execution_time"] else "н/д"
                response += f"{entity_label} {html.escape(str(specialist))}: {data['total_requests']} заявок, время {e_str}\n"
            response += f"\nВсего: {total_requests} заявок, время {total_exec_str}"
        await safe_edit_message_text(
            callback.bot,
            chat_id=callback.message.chat.id,
            message_id=callback.message.message_id,
            text=response, reply_markup=None
        )
        await safe_answer_callback(callback, "Статистика отображена")

    elif output_format == "excel":
        await _stats_excel(callback, stats, stats_type, headers, excel_sheet_title, excel_filename, period, start_date, end_date)


async def _stats_excel(
    callback: CallbackQuery,
    stats: list,
    stats_type: str,
    headers: list,
    sheet_title: str,
    filename: str,
    period: str,
    start_date: Optional[str],
    end_date: Optional[str]
) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append([sheet_title])
    ws.append([format_period(period, start_date, end_date)])
    ws.append([])
    ws.append(headers)

    if stats_type == "specialists":
        cur_spec = None
        spec_req = 0
        spec_time = 0.0
        for specialist, chat_id, count, exec_time in stats:
            try:
                chat = await callback.bot.get_chat(int(chat_id))
                chat_name = chat.title or f"Чат {chat_id}"
            except Exception:
                chat_name = f"Чат {chat_id}"
            if cur_spec != specialist:
                if cur_spec is not None:
                    ws.append(["Всего по специалисту", "", spec_req, _fmt_td(spec_time)])
                    ws.append([])
                ws.append([specialist, "", "", ""])
                cur_spec = specialist
                spec_req = 0
                spec_time = 0.0
            ws.append(["", chat_name, count, _fmt_td(exec_time) if exec_time else "н/д"])
            spec_req += int(count)
            spec_time += exec_time or 0
        if cur_spec is not None:
            ws.append(["Всего по специалисту", "", spec_req, _fmt_td(spec_time)])
            ws.append([])
    else:
        cur_chat = None
        chat_req = 0
        chat_exec = 0.0
        chat_wait = 0.0
        for chat_id, specialist, count, exec_time, wait_time in stats:
            try:
                chat = await callback.bot.get_chat(int(chat_id))
                chat_name = chat.title or f"Чат {chat_id}"
            except Exception:
                chat_name = f"Чат {chat_id}"
            if cur_chat != chat_id:
                if cur_chat is not None:
                    ws.append(["Всего по чату", "", chat_req, _fmt_td(chat_exec), _fmt_td(chat_wait)])
                    ws.append([])
                ws.append([chat_name, "", "", "", ""])
                cur_chat = chat_id
                chat_req = 0
                chat_exec = 0.0
                chat_wait = 0.0
            ws.append(["", specialist, count,
                        _fmt_td(exec_time) if exec_time else "н/д",
                        _fmt_td(wait_time) if wait_time else "н/д"])
            chat_req += int(count)
            chat_exec += exec_time or 0
            chat_wait += wait_time or 0
        if cur_chat is not None:
            ws.append(["Всего по чату", "", chat_req, _fmt_td(chat_exec), _fmt_td(chat_wait)])
            ws.append([])

    for i in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 20

    with NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
        tmp_path = tmp.name
    try:
        await asyncio.to_thread(save_workbook, wb, tmp_path)
        await callback.message.delete()
        await callback.message.answer_document(
            types.FSInputFile(tmp_path, filename=f"{filename}.xlsx")
        )
        await safe_answer_callback(callback, "Excel-файл отправлен")
    finally:
        os.unlink(tmp_path)


# ---------------------------------------------------------------------------
# my_stats: — личная статистика специалиста
# ---------------------------------------------------------------------------

@router.callback_query(F.data.startswith("my_stats:"))
async def cb_my_stats(callback: CallbackQuery, state: FSMContext) -> None:
    if not callback.message:
        await safe_answer_callback(callback, "Сообщение недоступно", show_alert=True)
        return
    roles = await check_user_role(callback.bot, callback.from_user.id, callback.message.chat.id, callback.message.chat.type)
    allowed = roles["is_admin_or_botadmin"] or roles["is_specialist"]
    if not allowed:
        await state.clear()
        await safe_answer_callback(callback, "⛔ Доступ запрещён", show_alert=True)
        return
    if not _valid_stats_data(callback.data):
        await safe_answer_callback(callback, "Некорректная кнопка статистики", show_alert=True)
        return
    parts = callback.data.split(":")
    action = parts[1]

    if action == "period":
        period = parts[2]
        if period == "custom":
            await state.update_data(stats_type="specialist_self")
            await safe_edit_message_text(
                callback.bot,
                chat_id=callback.message.chat.id,
                message_id=callback.message.message_id,
                text="📅 Введите начальную дату в формате ГГГГ-ММ-ДД (например, 2025-05-01):",
                reply_markup=None
            )
            await state.set_state(StatsStates.waiting_for_custom_start_date)
        else:
            keyboard = InlineKeyboardMarkup(inline_keyboard=[[
                InlineKeyboardButton(text="В Telegram", callback_data=f"my_stats:telegram:{period}"),
                InlineKeyboardButton(text="Скачать Excel", callback_data=f"my_stats:excel:{period}")
            ]])
            await safe_edit_message_text(
                callback.bot,
                chat_id=callback.message.chat.id,
                message_id=callback.message.message_id,
                text="📊 Выберите формат вывода вашей статистики:",
                reply_markup=keyboard
            )
        await safe_answer_callback(callback, "")

    elif action in ("telegram", "excel"):
        output_format = action
        period = parts[2]
        start_date = parts[3] if len(parts) > 3 else None
        end_date = parts[4] if len(parts) > 4 else None

        specialist_name = callback.from_user.full_name
        stats = await database.get_specialist_statistics(period, start_date, end_date, specialist_id=callback.from_user.id)
        title = f"📊 Ваша статистика ({html.escape(specialist_name)})"
        excel_filename = f"my_stats_{period}"

        if not stats:
            period_text = format_period(period, start_date, end_date)
            await safe_edit_message_text(
                callback.bot,
                chat_id=callback.message.chat.id,
                message_id=callback.message.message_id,
                text=f"{title} ({period_text})\n\nНет данных для выбранного периода.",
                reply_markup=None
            )
            await safe_answer_callback(callback, "Нет данных для отображения")
            return

        total_requests = sum(int(c) for _, _, c, _ in stats)
        total_exec = sum(t for _, _, _, t in stats if t)
        total_exec_str = _fmt_td(total_exec) if total_exec else "н/д"

        if output_format == "telegram":
            period_text = format_period(period, start_date, end_date)
            response = f"{title} ({period_text})\n\n"
            for _, chat_id, count, exec_time in stats:
                try:
                    chat = await callback.bot.get_chat(int(chat_id))
                    name = html.escape(chat.title or f"Чат {chat_id}")
                except Exception:
                    name = f"Чат {chat_id}"
                e_str = _fmt_td(exec_time) if exec_time else "н/д"
                response += f"💬 {name}: {count} заявок, время {e_str}\n"
            response += f"\nВсего: {total_requests} заявок, время {total_exec_str}"
            await safe_edit_message_text(
                callback.bot,
                chat_id=callback.message.chat.id,
                message_id=callback.message.message_id,
                text=response, reply_markup=None
            )
            await safe_answer_callback(callback, "Статистика отображена")

        elif output_format == "excel":
            headers = ["Чат", "Количество заявок", "Время исполнения"]
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "My Statistics"
            period_text = format_period(period, start_date, end_date)
            ws.append(["My Statistics"])
            ws.append([f"Специалист: {specialist_name}"])
            ws.append([period_text])
            ws.append([])
            ws.append(headers)
            for _, chat_id, count, exec_time in stats:
                try:
                    chat = await callback.bot.get_chat(int(chat_id))
                    chat_name = chat.title or f"Чат {chat_id}"
                except Exception:
                    chat_name = f"Чат {chat_id}"
                ws.append([chat_name, count, _fmt_td(exec_time) if exec_time else "н/д"])
            ws.append(["Всего", total_requests, _fmt_td(total_exec)])
            for i in range(1, len(headers) + 1):
                ws.column_dimensions[get_column_letter(i)].width = 20
            with NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
                tmp_path = tmp.name
            try:
                await asyncio.to_thread(save_workbook, wb, tmp_path)
                await callback.message.delete()
                await callback.message.answer_document(
                    types.FSInputFile(tmp_path, filename=f"{excel_filename}.xlsx")
                )
                await safe_answer_callback(callback, "Excel-файл отправлен")
            finally:
                os.unlink(tmp_path)
